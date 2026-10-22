"""Paso 2: clasificación ABC de los productos por valor de venta.

Ejecuta las consultas SQL de la carpeta sql/ y deja dos resultados en
resultados/: un Excel listo para trabajar y la curva ABC en imagen.
"""

import sqlite3
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

CARPETA = Path(__file__).parent
BASE = CARPETA / "datos" / "inventario.db"
SQL = CARPETA / "sql"
RESULTADOS = CARPETA / "resultados"

COLORES = {"A": "#1F3864", "B": "#2E5C8A", "C": "#A6B8D4"}


def clasificar():
    with sqlite3.connect(BASE) as con:
        con.executescript((SQL / "01_ventas_limpias.sql").read_text())
        return pd.read_sql((SQL / "02_abc.sql").read_text(), con)


def resumir(abc):
    resumen = abc.groupby("clase").agg(productos=("producto", "count"), importe=("importe", "sum"))
    resumen["pct_productos"] = resumen["productos"] / resumen["productos"].sum()
    resumen["pct_importe"] = resumen["importe"] / resumen["importe"].sum()
    return resumen.reset_index()


def guardar_excel(abc, resumen, ruta):
    productos = abc.rename(columns={
        "producto": "Código", "descripcion": "Descripción", "unidades": "Unidades netas",
        "pedidos": "Pedidos", "importe": "Venta neta (£)", "pct_acumulado": "% acumulado",
        "clase": "Clase",
    })
    tabla = resumen.rename(columns={
        "clase": "Clase", "productos": "Productos", "pct_productos": "% de productos",
        "importe": "Venta neta (£)", "pct_importe": "% de la venta",
    })[["Clase", "Productos", "% de productos", "Venta neta (£)", "% de la venta"]]

    formatos = {
        "Unidades netas": "#,##0", "Pedidos": "#,##0", "Productos": "#,##0",
        "Venta neta (£)": "#,##0", "% acumulado": "0.0%", "% de productos": "0.0%",
        "% de la venta": "0.0%",
    }
    with pd.ExcelWriter(ruta, engine="openpyxl") as libro:
        for nombre, df in (("Resumen", tabla), ("Productos", productos)):
            df.to_excel(libro, sheet_name=nombre, index=False)
            hoja = libro.sheets[nombre]
            for i, columna in enumerate(df.columns, start=1):
                letra = get_column_letter(i)
                celda = hoja[f"{letra}1"]
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="1F3864")
                ancho = max(len(columna), df[columna].astype(str).str.len().max())
                hoja.column_dimensions[letra].width = min(ancho + 3, 45)
                if columna in formatos:
                    for fila in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                        fila[0].number_format = formatos[columna]
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions


def dibujar_curva(abc, resumen, ruta):
    pct_productos = (abc.index + 1) / len(abc) * 100
    pct_venta = abc["pct_acumulado"] * 100

    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=150)
    ax.plot(pct_productos, pct_venta, color="#1F3864", linewidth=2.2)

    inicio = 0
    for _, fila in resumen.iterrows():
        fin = inicio + fila["pct_productos"] * 100
        ax.axvspan(inicio, fin, color=COLORES[fila["clase"]], alpha=0.12)
        ax.text((inicio + fin) / 2, 6, fila["clase"], ha="center", fontsize=20,
                fontweight="bold", color=COLORES[fila["clase"]])
        inicio = fin

    a = resumen.set_index("clase").loc["A"]
    ax.annotate(f"El {a.pct_productos * 100:.0f} % de los productos\nhace el {a.pct_importe * 100:.0f} % de la venta",
                xy=(a.pct_productos * 100, a.pct_importe * 100),
                xytext=(a.pct_productos * 100 + 12, 55), fontsize=11,
                arrowprops={"arrowstyle": "->", "color": "#444"})

    ax.set_xlim(0, 100)
    ax.set_ylim(0, 102)
    ax.set_xlabel("% de productos (de más a menos venta)")
    ax.set_ylabel("% acumulado de la venta neta")
    ax.set_title(f"Curva ABC · {len(abc):,} productos · dic 2010 – nov 2011".replace(",", "."),
                 loc="left", fontsize=13)
    ax.grid(alpha=0.3)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    fig.savefig(ruta)


def main():
    RESULTADOS.mkdir(exist_ok=True)
    abc = clasificar()
    resumen = resumir(abc)
    guardar_excel(abc, resumen, RESULTADOS / "clasificacion_abc.xlsx")
    dibujar_curva(abc, resumen, RESULTADOS / "curva_abc.png")
    print(resumen.to_string(index=False))


if __name__ == "__main__":
    main()
