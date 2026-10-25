"""Formato común de las hojas de Excel que generan los scripts."""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def dar_formato(hoja, df, formatos, ancho_maximo=45):
    """Cabecera azul, anchos ajustados, formatos numéricos, fila fija y filtros."""
    for i, columna in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        celda = hoja[f"{letra}1"]
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3864")
        celda.alignment = Alignment(vertical="center")
        ancho = max(len(columna), df[columna].astype(str).str.len().max())
        hoja.column_dimensions[letra].width = min(ancho + 3, ancho_maximo)
        if columna in formatos:
            for fila in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                fila[0].number_format = formatos[columna]
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
